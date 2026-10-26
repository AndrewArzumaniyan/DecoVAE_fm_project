"""Typeset the complete Markdown report as an editable Word report.

Uses the bundled document runtime. Numeric data are read, not recomputed. Figure
layouts are redrawn at print size from the same saved intervals. The renderer
exports the PDF; a second pass fills the static, linked table of contents.
"""
import argparse
import json
import math
import re
import subprocess
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


RUNTIME=Path('/Users/andrew/.cache/codex-runtimes/codex-primary-runtime/dependencies')
POPPLER=RUNTIME/'bin/override/pdftoppm'
BLACK='000000'
BLUE='173C55'
TITLE_MAP={
    1:'Основные выводы проекта',2:'Постановка и протокол',
    3:'Архитектуры и функции потерь',4:'Проверка исходной декомпозиции',
    5:'Результаты на реальных данных',6:'Сезонная информация и структура разложения',
    7:'Синтетическая лаборатория и границы переноса',8:'Разнообразие данных и масштаб моделей',
    9:'Калибровка и обучение по смеси',10:'Проверка гипотезы о шумовой ветке',
    11:'Ошибки оценщика и регуляризация',12:'Согласованное внешнее сравнение',
    13:'Ответы на ревью',14:'Завершённые проверки и задачи к защите',15:'Данные и воспроизводимость'}


def readable(text):
    text=text.replace('Презентация этим Markdown-файлом не создаётся.','Файл презентации в состав отчёта не входит.')
    text=text.replace('Это исходные материалы, а не инструкции к запуску; неточности старых пересказов уточняются по коду и метрикам.',
                      'Неточности старых пересказов уточняются по коду и метрикам.')
    return text.translate(str.maketrans({'\u2011':'-','\u2013':'-','\u2014':'-','\u2212':'-'}))


def heading(text):
    text=re.sub(r'^\d+(?:\.\d+)*\.?\s*','',text)
    text=text.replace('+',' и ').replace('/',' и ')
    return ' '.join(re.sub(r'[^\w\s]',' ',text).replace('_',' ').split())


def make_print_figures(evidence, dest):
    dest.mkdir(parents=True,exist_ok=True)
    pdfmetrics.registerFont(TTFont('ReportArial','/System/Library/Fonts/Supplemental/Arial.ttf'))
    pdfmetrics.registerFont(TTFont('ReportArialBold','/System/Library/Fonts/Supplemental/Arial Bold.ttf'))
    W=511.2
    def txt(c,x,y,s,size=9,align='left',bold=False):
        c.setFillColorRGB(0,0,0);c.setFont('ReportArialBold' if bold else 'ReportArial',size)
        fn={'left':c.drawString,'right':c.drawRightString,'center':c.drawCentredString}[align]
        fn(x,y,readable(s))
    def graph(c,box,points,limits,ticks,reference=None,log=False):
        x0,y0,x1,y1=box
        def X(x):
            f=math.log10 if log else float
            return x0+(f(x)-f(limits[0]))/(f(limits[1])-f(limits[0]))*(x1-x0)
        c.setLineWidth(.55);c.setStrokeColorRGB(.25,.25,.25);c.line(x0,y0,x1,y0)
        for val,label in ticks:
            x=X(val);c.line(x,y0,x,y0-4);txt(c,x,y0-17,label,9,'center')
        if reference is not None:
            c.saveState();c.setDash(2,3);c.setStrokeColorRGB(.55,.55,.55)
            c.line(X(reference),y0,X(reference),y1);c.restoreState()
        for i,(label,point,low,high,color) in enumerate(points):
            y=y1-(i+.5)*(y1-y0)/len(points)
            txt(c,x0-10,y-3,label,9.5,'right')
            c.setStrokeColorRGB(*color);c.setFillColorRGB(*color);c.setLineWidth(1.15)
            c.line(X(low),y,X(high),y);c.line(X(low),y-4,X(low),y+4);c.line(X(high),y-4,X(high),y+4)
            c.circle(X(point),y,3,stroke=0,fill=1)
    def finish(c,name):
        c.save()
        subprocess.run([str(POPPLER),'-r','240','-singlefile','-png',str(dest/(name+'.pdf')),str(dest/name)],check=True,
                       stdout=subprocess.DEVNULL,stderr=subprocess.PIPE)
    colors={'Dmc1Zn':(.68,.28,.18),'Dmc1B4':(.14,.42,.60)}
    h=pd.read_csv(evidence/'h1_comparisons.csv')
    name='h1_five_seed_intervals';c=canvas.Canvas(str(dest/(name+'.pdf')),pagesize=(W,432))
    txt(c,W/2,412,'Пять сидов и парные 95% доверительные интервалы',11,'center',True)
    for box,tests,title,lim,ticks,ty in [
        ((225,251,490,380),[('real','ncrps','Реальные данные'),('strong','crps_h1','Сильная волатильность 1-32')],
         'Ошибка прогноза CRPS',(-1.2,1.35),[-1,-.5,0,.5,1],391),
        ((225,56,490,185),[('strong','lv_err','Сильная волатильность'),('base','lv_err','Базовая волатильность')],
         'Ошибка логарифма дисперсии',(-8.5,14),[-5,0,5,10],196)]:
        points=[]
        for profile,metric,label in tests:
            for arm in ['Dmc1Zn','Dmc1B4']:
                row=h[(h.profile==profile)&(h.metric==metric)&(h.arm==arm)].iloc[0]
                points.append((label+' '+('Zn' if arm.endswith('Zn') else 'B4'),100*(row.effect-1),
                               100*(row.ci_low-1),100*(row.ci_high-1),colors[arm]))
        txt(c,350,ty,title,11,'center',True)
        graph(c,box,points,lim,[(v,f'{v:g}'.replace('.',',')) for v in ticks],reference=0)
        txt(c,350,box[1]-34,'Изменение относительно PD, %',9,'center')
    txt(c,W/2,10,'Меньше лучше  Контрасты исследовательские без поправки на множественность',8.5,'center')
    finish(c,name)

    d=pd.read_csv(evidence/'component_intervals.csv')
    name='component_recovery_intervals';c=canvas.Canvas(str(dest/(name+'.pdf')),pagesize=(W,242))
    txt(c,W/2,221,'Восстановление сезонности на контексте и горизонте',11,'center',True)
    for x0,x1,region,label in [(47,241,'context','Контекст'),(307,501,'horizon','Горизонт')]:
        points=[]
        for model,col in [('A',(.15,.42,.61)),('Dmc1',(.85,.42,.10))]:
            row=d[(d.dataset=='slow_1024_plus_24_sigma0.3')&(d.model==model)&(d.metric=='seasonal_'+region+'_mse')].iloc[0]
            points.append(('A' if model=='A' else 'PD',row['mean'],row.ci_low,row.ci_high,col))
        txt(c,(x0+x1)/2,197,label,10.5,'center',True)
        graph(c,(x0,67,x1,181),points,(.0003,.65),[(.001,'0,001'),(.01,'0,01'),(.1,'0,1')],log=True)
        txt(c,(x0+x1)/2,32,'MSE / дисперсия контекста',9,'center')
    txt(c,W/2,10,'Волна 1024 и цикл 24  σ = 0,3  Три сида  Логарифмическая шкала MSE',9,'center')
    finish(c,name)

    name='loso_intervals';c=canvas.Canvas(str(dest/(name+'.pdf')),pagesize=(W,157))
    txt(c,W/2,139,'LOSO по 15 источникам',11,'center',True)
    a=pd.read_csv(evidence/'loso_comparisons.csv');labels={'Dmc1/N6mc1':'PD / Transformer','N6s/N6mc1':'N6s / Transformer','Dmc1/N6s':'PD / N6s'}
    points=[(labels[row.comparison],row.ratio,row.ci_low,row.ci_high,(.13,.40,.58)) for row in a.itertuples()]
    graph(c,(130,42,492,122),points,(.81,1.05),[(v,f'{v:.2f}'.replace('.',',')) for v in [.85,.90,.95,1,1.05]],reference=1)
    txt(c,310,9,'Отношение CRPS и 95% интервал по источникам',9.5,'center')
    finish(c,name)

    name='weak_seasonality_leakage';c=canvas.Canvas(str(dest/(name+'.pdf')),pagesize=(W,282))
    txt(c,W/2,262,'Слабый цикл и зависимость от уровня шума',11,'center',True)
    x0,y0,x1,y1=65,55,492,229
    X=lambda v:x0+(v-.08)/(.55)*(x1-x0)
    Y=lambda v:y0+(v+.05)/1.16*(y1-y0)
    c.setStrokeColorRGB(.25,.25,.25);c.setLineWidth(.55);c.line(x0,y0,x1,y0);c.line(x0,y0,x0,y1)
    for v in [.1,.3,.6]:
        c.line(X(v),y0,X(v),y0-4);txt(c,X(v),y0-17,str(v).replace('.',','),9,'center')
    for v in [0,.25,.5,.75,1]:
        txt(c,x0-8,Y(v)-3,f'{v:g}'.replace('.',','),9,'right')
    for v in [0,1]:
        c.saveState();c.setDash(2,3);c.setStrokeColorRGB(.6,.6,.6);c.line(x0,Y(v),x1,Y(v));c.restoreState()
    for model,label,color,lx in [('A','Исходная A',(.13,.40,.58),288),('Dmc1','PD',(.85,.42,.10),418)]:
        part=d[d.dataset.str.startswith('weak_seasonality_sigma')&d.model.eq(model)&d.metric.eq('seasonal_in_trend_gain')].copy()
        part['sigma']=part.dataset.str.split('sigma').str[-1].astype(float);part=part.sort_values('sigma')
        c.setStrokeColorRGB(*color);c.setFillColorRGB(*color);c.setLineWidth(1.1)
        last=None
        for row in part.itertuples():
            x,y=X(row.sigma),Y(row.mean)
            if last:c.line(*last,x,y)
            c.line(x,Y(row.ci_low),x,Y(row.ci_high));c.line(x-3,Y(row.ci_low),x+3,Y(row.ci_low));c.line(x-3,Y(row.ci_high),x+3,Y(row.ci_high))
            c.circle(x,y,3,fill=1,stroke=0);last=(x,y)
        c.circle(lx,244,3,fill=1,stroke=0);txt(c,lx+8,241,label,9)
    txt(c,(x0+x1)/2,17,'Стандартное отклонение шума σ',9.5,'center')
    c.saveState();c.translate(15,142);c.rotate(90);txt(c,0,0,'Сезонная проекция ошибки тренда',9,'center');c.restoreState()
    finish(c,name)


def xmlfont(run,name='Arial',size=None):
    run.font.name=name
    rp=run._element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
    for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),name)
    if size:run.font.size=Pt(size)


def link(paragraph,label,target,anchor=False,size=None):
    h=OxmlElement('w:hyperlink')
    if anchor:h.set(qn('w:anchor'),target);h.set(qn('w:history'),'1')
    else:
        from docx.opc.constants import RELATIONSHIP_TYPE
        h.set(qn('r:id'),paragraph.part.relate_to(target,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),BLACK if anchor else '225B7C');rp.append(color)
    if not anchor:
        u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u)
    if size:
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(round(size*2)));rp.append(sz)
    r.append(rp);t=OxmlElement('w:t');t.text=readable(label);r.append(t);h.append(r);paragraph._p.append(h)


INLINE=re.compile(r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<!\*)\*[^*\n]+\*(?!\*))')


def inline(p,text,bold=False,size=None):
    for token in INLINE.split(text):
        if not token:continue
        if token.startswith('**') and token.endswith('**'):inline(p,token[2:-2],True,size)
        elif token.startswith('*') and token.endswith('*'):
            r=p.add_run(readable(token[1:-1]));r.bold=bold;r.italic=True
            if size:r.font.size=Pt(size)
        elif token.startswith('`') and token.endswith('`'):
            r=p.add_run(readable(token[1:-1]));r.bold=bold;xmlfont(r,'Courier New',min(size or 10,9.5))
        elif token.startswith('[') and '](' in token:
            label,target=token[1:].split('](',1);link(p,label,target[:-1],size=size)
        else:
            r=p.add_run(readable(token));r.bold=bold
            if size:r.font.size=Pt(size)


def math_run(text):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r


def sub(base,lower):
    node=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(math_run(base));s=OxmlElement('m:sub');s.append(math_run(lower));node.extend([e,s]);return node


def equation(doc,kind):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(9)
    outer=OxmlElement('m:oMathPara');m=OxmlElement('m:oMath')
    if kind=='loss':
        nodes=[math_run('L = 0.2 · '),sub('NLL','recon'),math_run(' + '),sub('NLL','mix'),math_run(' + afCRPS + 0.01 · KL')]
    elif kind=='quantile':
        nodes=[sub('qCRPS','9'),math_run(' = 2 · '),sub('mean','q ∈ {0.1,…,0.9}'),math_run('('),sub('pinball','q'),math_run(')')]
    else:
        frac=OxmlElement('m:f');num=OxmlElement('m:num');num.append(sub('R','d'));den=OxmlElement('m:den');den.append(sub('R','d,SN'));frac.extend([num,den])
        nodes=[math_run('skill = exp('),sub('mean','d'),math_run('[log('),frac,math_run(')])')]
    for node in nodes:m.append(node)
    outer.append(m);p._p.append(outer)
    if kind=='skill':
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8)
        inline(p,'R_d - средняя нормированная CRPS модели на источнике d; R_d,SN - та же ошибка seasonal-naive.',size=9.5)


def widths(headers):
    n=len(headers)
    if n==2:return [.37,.63]
    if n==3:
        if headers[0]=='ID в коде':return [.15,.25,.60]
        if headers[0]=='Метрика':return [.40,.30,.30]
        return [.43,.22,.35]
    if n==4:
        if headers[0]=='Горизонт':return [.12,.2933,.2933,.2934]
        if headers[0]=='Модель' and headers[1]=='Сиды':return [.43,.09,.24,.24]
        if headers[0]=='Модель' and 'MSE' in headers[1]:return [.15,.34,.34,.17]
        return [.46,.18,.18,.18]
    if n==5:
        if headers[0].startswith('Данные'):return [.34,.10,.23,.10,.23]
        if headers[0]=='Сравнение':return [.34,.14,.26,.13,.13]
        if headers[0]=='Модель' and 'Сиды' in headers[1]:return [.32,.15,.15,.24,.14]
        if headers[0]=='Модель':return [.30,.175,.175,.175,.175]
        return [.36,.16,.16,.16,.16]
    return [.30]+[.70/(n-1)]*(n-1)


def table(doc,rows,review=False):
    if review:
        for index,row in enumerate(rows[1:],1):
            p=doc.add_paragraph(style='Heading 2');p.add_run(f'{index} '+heading(row[0]))
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            inline(p,'**Ответ.** '+row[1]+('' if row[1].endswith(('.',':',';')) else '.'))
            p=doc.add_paragraph();inline(p,'**Ограничения.** '+row[2]+('' if row[2].endswith(('.',':',';')) else '.'))
        return
    t=doc.add_table(rows=1,cols=len(rows[0]));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    ws=widths(rows[0]);total=7.05
    for col,width in zip(t.columns,ws):col.width=Inches(total*width)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    margin=OxmlElement('w:tblCellMar')
    for side,v in [('top',85),('bottom',85),('left',85),('right',85)]:
        x=OxmlElement('w:'+side);x.set(qn('w:w'),str(v));x.set(qn('w:type'),'dxa');margin.append(x)
    pr.append(margin)
    for index,values in enumerate(rows):
        row=t.rows[0] if index==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if index==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,text) in enumerate(zip(row.cells,values)):
            cell.width=Inches(total*ws[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.08
            p.paragraph_format.keep_with_next=len(rows)<=8 and index<len(rows)-1
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(values)<=3 else WD_ALIGN_PARAGRAPH.CENTER
            inline(p,text,size=9.5 if index==0 else 10)
            if index==0:
                for run in p.runs:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),BLUE if index==0 else ('F3F6F8' if index%2==0 else 'FFFFFF'))
            cell._tc.get_or_add_tcPr().append(shd)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
    p.paragraph_format.line_spacing=Pt(3);p.add_run().font.size=Pt(3)


def build(source,out,qa,page_map=None):
    doc=Document();sec=doc.sections[0]
    for st in doc.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
    sec.footer_distance=Inches(.28);sec.different_first_page_header_footer=True
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11);normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.widow_control=True
    for name,size in [('Title',26),('Subtitle',18),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
        st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.font.italic=False
        st.font.bold=name.startswith('Heading')
        st.paragraph_format.space_before=Pt(15 if name=='Heading 1' else 10)
        st.paragraph_format.space_after=Pt(7);st.paragraph_format.keep_with_next=True
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run('Страница ');r.font.size=Pt(9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE')
    fr=OxmlElement('w:r');fp=OxmlElement('w:rPr');fs=OxmlElement('w:sz');fs.set(qn('w:val'),'18');fp.append(fs);fr.append(fp)
    ft=OxmlElement('w:t');ft.text='1';fr.append(ft);fld.append(fr);footer._p.append(fld)
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
    doc.core_properties.title='Полный отчёт по проекту DecoVAE FM'
    doc.core_properties.subject='Эксперименты декомпозиция вероятностный прогноз и ответы на ревью'
    doc.core_properties.author='';doc.core_properties.last_modified_by=''

    p=doc.add_paragraph(style='Title');p.paragraph_format.space_before=Pt(48);p.add_run('DecoVAE FM')
    doc.add_paragraph('Полный отчёт по экспериментам к защите',style='Subtitle')
    p=doc.add_paragraph('25 октября 2026 года');p.paragraph_format.space_after=Pt(28)
    p=doc.add_paragraph();inline(p,'**Основной результат.** Вычисленная сезонная информация улучшает перенос на новые сезонные источники в малом режиме обучения. Необходимость жёсткой декомпозиции и польза дополнительного шумового латента не установлены.')
    p=doc.add_paragraph('В отчёте разобраны архитектуры и функции потерь, контролируемые абляции, перенос между источниками, синтетические проверки компонент, калибровка, пятисидовый H1 и внешнее сравнение с Chronos. Результаты сопровождаются условиями экспериментов и ограничениями выводов.')
    lines=source.read_text().splitlines()
    for line in lines:
        if line.startswith(('**Дата:**','**Статус:**','**Основание:**')):
            p=doc.add_paragraph();inline(p,line.strip(),size=9.5)
    p=doc.add_paragraph('Качество прогноза • Декомпозиция • Неопределённость');p.paragraph_format.space_before=Pt(22)
    p.runs[0].font.size=Pt(11)
    doc.add_page_break()
    doc.add_paragraph('Содержание',style='Title')
    chapters=[]
    for n,title in TITLE_MAP.items():
        label=f'{n} {title}';chapters.append(dict(number=n,title=label,anchor=f'chapter_{n:02d}'))
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(9)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.05),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        link(p,label,f'chapter_{n:02d}',anchor=True)
        p.add_run('\t'+str((page_map or {}).get(str(n),'')))
    doc.add_page_break()
    i=0;figures=[];table_count=0;body_paragraphs=0;main=False;bookmark_id=1
    while i<len(lines):
        line=lines[i].strip()
        if line.startswith('## '):
            main=True;n=int(re.match(r'##\s+(\d+)',line)[1]);p=doc.add_paragraph(style='Heading 1')
            b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(bookmark_id));b.set(qn('w:name'),f'chapter_{n:02d}');p._p.append(b)
            p.add_run(f'{n} '+TITLE_MAP[n]);e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(bookmark_id));p._p.append(e);bookmark_id+=1
            i+=1;continue
        if not main or not line:i+=1;continue
        if line.startswith('### '):doc.add_paragraph(heading(line[4:]),style='Heading 2');i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                vals=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',v) for v in vals):rows.append(vals)
                i+=1
            table(doc,rows,review=rows[0][0]=='Замечание / вопрос');table_count+=1;continue
        match=re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)',line)
        if match:
            label,target=match.groups();original=(source.parent/target).resolve();image=qa/'print_figures'/original.name
            if not image.exists():image=original
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(image),width=Inches(7.05))
            for obj in p._p.iter(qn('wp:docPr')):obj.set('descr',label)
            cap=doc.add_paragraph();cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after=Pt(10);inline(cap,f'Рисунок {len(figures)+1}. '+label,size=9.5)
            figures.append(dict(source=str(original),embedded=str(image),caption=label));i+=1;continue
        if line.startswith('`skill(model)'):equation(doc,'skill');i+=1;continue
        if line.startswith('`L ='):equation(doc,'loss');i+=1;continue
        if line.startswith('`qCRPS9 ='):equation(doc,'quantile');i+=1;continue
        if re.match(r'^\d+\.\s',line):
            p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.22);p.paragraph_format.first_line_indent=Inches(-.22)
            inline(p,line)
        elif line.startswith('- '):
            text=line[2:];text=re.sub(r'^\[x\]\s*','Выполнено - ',text);text=re.sub(r'^\[ \]\s*','Осталось - ',text)
            p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.16);p.paragraph_format.first_line_indent=Inches(-.16)
            inline(p,'• '+text)
        else:
            p=doc.add_paragraph();inline(p,line)
        body_paragraphs+=1;i+=1
    out.parent.mkdir(parents=True,exist_ok=True);doc.save(out)
    (qa/'document_manifest.json').write_text(json.dumps(dict(source=str(source),output=str(out),chapters=chapters,
          markdown_tables=table_count,word_tables=len(doc.tables),figures=figures,body_paragraphs=body_paragraphs,
          page_map=page_map or {}),ensure_ascii=False,indent=2)+'\n')
    print('Created',out,'tables',len(doc.tables),'figures',len(figures),'chapters',len(chapters))


def main():
    p=argparse.ArgumentParser();p.add_argument('--source',type=Path,required=True);p.add_argument('--out',type=Path,required=True)
    p.add_argument('--qa',type=Path,required=True);p.add_argument('--page-map',type=Path)
    p.add_argument('--reuse-figures',action='store_true');a=p.parse_args();a.qa.mkdir(parents=True,exist_ok=True)
    if not a.reuse_figures:make_print_figures(a.source.parent/'evidence/20261025',a.qa/'print_figures')
    build(a.source.resolve(),a.out.resolve(),a.qa.resolve(),json.loads(a.page_map.read_text()) if a.page_map else None)


if __name__=='__main__':main()
